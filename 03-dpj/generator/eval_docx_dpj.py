#!/usr/bin/env python3
"""Ekstrak teks DOCX (python-docx) + evaluasi GEPA evaluator --doc untuk semua DPJ.
Output: skor per dokumen + JSON hasil. Exit 1 bila ada skor < 0.8."""
import json, os, subprocess, sys
import docx

DOCS = [
    ("DPJ-1", "/home/z/my-project/download/proposal-pengajuan-judul/DPJ-1_Proposal_Judul_Kandang_Ayam_Broiler.docx"),
    ("DPJ-2", "/home/z/my-project/download/proposal-pengajuan-judul/DPJ-2_Proposal_Judul_Ruang_Kelas_SLB.docx"),
    ("DPJ-3", "/home/z/my-project/download/proposal-pengajuan-judul/DPJ-3_Proposal_Judul_Kamar_Asrama_SLB.docx"),
    ("DPJ-4", "/home/z/my-project/download/proposal-pengajuan-judul/DPJ-4_Proposal_Judul_Translasi_Isyarat_SIBI_BISINDO.docx"),
]
EVAL = "/home/z/my-project/skills/proposal-writing/scripts/evaluator.py"
OUT = "/home/z/my-project/tool-results/dpj_v2_eval"
os.makedirs(OUT, exist_ok=True)

results = {}
ok = True
for name, path in DOCS:
    d = docx.Document(path)
    parts = []
    for p in d.paragraphs:
        parts.append(p.text)
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    txt = "\n".join(parts)
    txtpath = os.path.join(OUT, name + ".txt")
    with open(txtpath, "w", encoding="utf-8") as f:
        f.write(txt)
    r = subprocess.run(["python3", EVAL, "--doc", txtpath, "--json-out", os.path.join(OUT, name + ".json")],
                       capture_output=True, text=True)
    info = json.loads(open(os.path.join(OUT, name + ".json")).read())
    results[name] = {"score": info["score"], "coverage": info["coverage"], "failed": info["failed"]}
    status = "OK " if info["score"] >= 0.8 else "LOW"
    if info["score"] < 0.8:
        ok = False
    print(f"{status} {name}: score={info['score']} cov={info['coverage']}")
    for f_ in info["failed"]:
        print("   FAIL:", f_.get("check"), "->", str(f_.get("suggestion"))[:120])

with open(os.path.join(OUT, "summary.json"), "w") as f:
    json.dump(results, f, indent=1)
sys.exit(0 if ok else 1)
